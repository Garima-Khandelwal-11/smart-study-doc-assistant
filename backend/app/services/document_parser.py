import pdfplumber
from docx import Document
from pptx import Presentation


def load_pdf_text(path):
    text = ""
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text += page.extract_text() + "\n"
    return text


def extract_text_from_docx(path):
    document = Document(path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def extract_text_from_pptx(path):
    presentation = Presentation(path)
    text_runs = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text_runs.append(shape.text_frame.text)
    return "\n".join(text_runs)


def extract_text_from_txt(path):
    return path.read().decode("utf-8")
